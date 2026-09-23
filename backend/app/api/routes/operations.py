import csv
import io
import json
import uuid
from datetime import UTC, date, datetime, timedelta
from itertools import islice

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import Response
from geoalchemy2.functions import ST_MakePoint, ST_SetSRID
from openpyxl import load_workbook
from sqlalchemy import delete, func, or_, select, update
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.audit import add_audit
from app.corridor_service import waypoint_model
from app.database import get_db
from app.dependencies import admin_user, csrf_protect, current_user, editor_user
from app.models import (
    AlertAcknowledgement,
    AppSetting,
    AuditLog,
    BackgroundJob,
    Corridor,
    CorridorWaypoint,
    CountryGateway,
    CustomsPost,
    DataImport,
    PostDailyMetric,
    SavedFilter,
    TransitDeclaration,
    User,
)
from app.schemas import SavedFilterCreate
from app.jobs import run_route_rebuild
from app.query_utils import contains_pattern

router = APIRouter(tags=["operations"])
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


async def csv_rows(file: UploadFile) -> list[dict[str, str]]:
    raw = await file.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="CSV fayl 5 MB dan katta bo'lmasin")
    if (file.filename or "").lower().endswith(".xlsx"):
        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator)]
            rows = []
            for values in iterator:
                row = {}
                for key, value in zip(headers, values):
                    row[key] = value.isoformat() if isinstance(value, (date, datetime)) else str(value if value is not None else "").strip()
                rows.append(row)
                if len(rows) > MAX_IMPORT_ROWS:
                    break
        except (OSError, ValueError, StopIteration) as exc:
            raise HTTPException(status_code=422, detail="Excel faylini o'qib bo'lmadi") from exc
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="CSV UTF-8 formatida bo'lishi kerak") from exc
        rows = [{str(key).strip(): str(value or "").strip() for key, value in row.items()} for row in islice(csv.DictReader(io.StringIO(text)), MAX_IMPORT_ROWS + 1)]
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=422, detail=f"Bir importda ko'pi bilan {MAX_IMPORT_ROWS} qator ruxsat etiladi")
    return rows


def parse_date(value: str, field: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{field} YYYY-MM-DD formatida bo'lishi kerak") from exc


def number(value: str, integer: bool = True) -> int | float:
    if not value:
        return 0
    return int(float(value)) if integer else float(value)


def normalize_import_row(import_type: str, row: dict[str, str]) -> dict:
    if import_type == "declarations":
        required = ["declaration_no", "declaration_date", "origin_country_code", "destination_country_code", "entry_post_code", "exit_post_code"]
        missing = [field for field in required if not row.get(field)]
        if missing:
            raise ValueError("Majburiy ustunlar bo'sh: " + ", ".join(missing))
        return {
            "declaration_no": row["declaration_no"], "declaration_date": parse_date(row["declaration_date"], "declaration_date"),
            "origin_country_code": row["origin_country_code"].upper(), "destination_country_code": row["destination_country_code"].upper(),
            "entry_post_code": row["entry_post_code"], "exit_post_code": row["exit_post_code"], "source_system": row.get("source_system") or "CSV",
            "vehicle_no": row.get("vehicle_no") or None, "carrier_name": row.get("carrier_name") or None, "state": row.get("state") or None,
        }
    if import_type == "posts":
        required = ["post_code", "post_name", "post_type"]
        missing = [field for field in required if not row.get(field)]
        if missing:
            raise ValueError("Majburiy ustunlar bo'sh: " + ", ".join(missing))
        return {
            "post_code": row["post_code"], "post_name": row["post_name"], "post_type": row["post_type"].upper(),
            "post_category": (row.get("post_category") or "UNASSIGNED").upper(), "region": row.get("region") or None,
            "neighbor_country_code": row.get("neighbor_country_code", "").upper() or None,
            "latitude": float(row["latitude"]) if row.get("latitude") else None, "longitude": float(row["longitude"]) if row.get("longitude") else None,
            "is_active": row.get("is_active", "true").lower() not in {"0", "false", "no"},
        }
    if import_type == "metrics":
        required = ["post_code", "post_type", "metric_date"]
        missing = [field for field in required if not row.get(field)]
        if missing:
            raise ValueError("Majburiy ustunlar bo'sh: " + ", ".join(missing))
        values = {"post_code": row["post_code"], "post_type": row["post_type"].upper(), "metric_date": parse_date(row["metric_date"], "metric_date")}
        float_fields = {"narcotics_kg", "customs_payments", "additional_customs_payments"}
        metric_fields = [column.name for column in PostDailyMetric.__table__.columns if column.name not in {"id", "post_code", "post_type", "metric_date", "created_at", "updated_at"}]
        values.update({field: number(row.get(field, ""), field not in float_fields) for field in metric_fields})
        return values
    raise HTTPException(status_code=422, detail="import_type declarations, posts yoki metrics bo'lishi kerak")


def validate_rows(import_type: str, rows: list[dict[str, str]]) -> tuple[list[dict], list[dict]]:
    valid: list[dict] = []
    errors: list[dict] = []
    for index, row in enumerate(rows, start=2):
        try:
            valid.append(normalize_import_row(import_type, row))
        except (ValueError, TypeError) as exc:
            errors.append({"row": index, "message": str(exc)[:300]})
    return valid, errors


@router.post("/imports/preview", dependencies=[Depends(csrf_protect)])
async def preview_import(import_type: str = Form(...), file: UploadFile = File(...), _: User = Depends(editor_user)) -> dict:
    rows = await csv_rows(file)
    valid, errors = validate_rows(import_type, rows)
    return {"filename": file.filename, "import_type": import_type, "rows_total": len(rows), "valid": len(valid), "rejected": len(errors), "preview": jsonable_encoder(valid[:20]), "errors": errors[:100]}


@router.post("/imports/commit", dependencies=[Depends(csrf_protect)])
async def commit_import(request: Request, import_type: str = Form(...), file: UploadFile = File(...), db: AsyncSession = Depends(get_db), user: User = Depends(editor_user)) -> dict:
    rows = await csv_rows(file)
    valid, errors = validate_rows(import_type, rows)
    if not valid:
        raise HTTPException(status_code=422, detail="Import qilinadigan yaroqli qator topilmadi")
    if import_type == "declarations":
        statement = insert(TransitDeclaration).values(valid)
        statement = statement.on_conflict_do_update(index_elements=[TransitDeclaration.declaration_no], set_={key: getattr(statement.excluded, key) for key in valid[0] if key != "declaration_no"})
    elif import_type == "posts":
        statement = insert(CustomsPost).values(valid)
        statement = statement.on_conflict_do_update(index_elements=[CustomsPost.post_code], set_={key: getattr(statement.excluded, key) for key in valid[0] if key != "post_code"})
    else:
        statement = insert(PostDailyMetric).values(valid)
        statement = statement.on_conflict_do_update(index_elements=[PostDailyMetric.post_code, PostDailyMetric.metric_date], set_={key: getattr(statement.excluded, key) for key in valid[0] if key not in {"post_code", "metric_date"}})
    await db.execute(statement)
    if import_type == "posts":
        imported_codes = [row["post_code"] for row in valid]
        await db.execute(update(CustomsPost).where(CustomsPost.post_code.in_(imported_codes), CustomsPost.latitude.is_not(None), CustomsPost.longitude.is_not(None)).values(location=ST_SetSRID(ST_MakePoint(CustomsPost.longitude, CustomsPost.latitude), 4326)))
    record = DataImport(import_type=import_type, filename=(file.filename or "import.csv")[:255], status="COMPLETED" if not errors else "COMPLETED_WITH_ERRORS", rows_total=len(rows), rows_imported=len(valid), rows_rejected=len(errors), errors=errors[:500], created_by=user.id)
    db.add(record)
    await db.flush()
    await add_audit(db, request, user, "IMPORT", import_type, str(record.id), after={"filename": record.filename, "imported": len(valid), "rejected": len(errors)})
    await db.commit()
    return {"id": str(record.id), "rows_total": len(rows), "rows_imported": len(valid), "rows_rejected": len(errors), "errors": errors[:100]}


@router.get("/imports")
async def list_imports(db: AsyncSession = Depends(get_db), _: User = Depends(editor_user)) -> dict:
    rows = (await db.scalars(select(DataImport).order_by(DataImport.created_at.desc()).limit(100))).all()
    return {"items": [{"id": str(row.id), "import_type": row.import_type, "filename": row.filename, "status": row.status, "rows_total": row.rows_total, "rows_imported": row.rows_imported, "rows_rejected": row.rows_rejected, "errors": row.errors, "created_at": row.created_at} for row in rows]}


@router.get("/data-status")
async def data_status(db: AsyncSession = Depends(get_db), _: User = Depends(current_user)) -> dict:
    latest_import = await db.scalar(select(func.max(DataImport.created_at)))
    latest_metric = await db.scalar(select(func.max(PostDailyMetric.metric_date)))
    latest_declaration = await db.scalar(select(func.max(TransitDeclaration.declaration_date)))
    counts = (await db.execute(select(
        select(func.count()).select_from(CustomsPost).scalar_subquery(),
        select(func.count()).select_from(Corridor).scalar_subquery(),
        select(func.count()).select_from(TransitDeclaration).scalar_subquery(),
        select(func.count()).select_from(PostDailyMetric).scalar_subquery(),
    ))).one()
    return {"latest_import": latest_import, "latest_metric_date": latest_metric, "latest_declaration_date": latest_declaration, "counts": {"posts": counts[0], "corridors": counts[1], "declarations": counts[2], "metrics": counts[3]}}


@router.get("/alerts")
async def alerts(db: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> dict:
    since = datetime.now(UTC) - timedelta(hours=1)
    review, unlocated, failed_logins = (await db.execute(select(
        select(func.count()).select_from(Corridor).where(Corridor.route_needs_review.is_(True)).scalar_subquery(),
        select(func.count()).select_from(CustomsPost).where(CustomsPost.is_active.is_(True), or_(CustomsPost.latitude.is_(None), CustomsPost.longitude.is_(None))).scalar_subquery(),
        select(func.count()).select_from(AuditLog).where(AuditLog.action == "LOGIN_FAILED", AuditLog.created_at >= since).scalar_subquery(),
    ))).one()
    latest_metric = await db.scalar(select(func.max(PostDailyMetric.metric_date)))
    candidates = []
    if review:
        candidates.append({"key": "routes-review", "level": "warning", "title": "Route tekshiruvi", "message": f"{review} ta koridor tekshiruv talab qiladi"})
    if unlocated:
        candidates.append({"key": "posts-unlocated", "level": "warning", "title": "Lokatsiyasiz postlar", "message": f"{unlocated} ta faol post koordinatasiz"})
    if failed_logins >= 5:
        candidates.append({"key": f"failed-logins-{datetime.now(UTC).date()}", "level": "danger", "title": "Login urinishlari", "message": f"Oxirgi soatda {failed_logins} ta muvaffaqiyatsiz urinish"})
    if latest_metric and latest_metric < date.today() - timedelta(days=45):
        candidates.append({"key": "stale-metrics", "level": "warning", "title": "Statistika eskirgan", "message": f"Oxirgi ko'rsatkich sanasi: {latest_metric}"})
    acknowledged = set((await db.scalars(select(AlertAcknowledgement.alert_key).where(AlertAcknowledgement.user_id == user.id))).all())
    return {"items": [{**item, "acknowledged": item["key"] in acknowledged} for item in candidates]}


@router.post("/alerts/{alert_key}/acknowledge", dependencies=[Depends(csrf_protect)])
async def acknowledge_alert(alert_key: str, db: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if not await db.scalar(select(AlertAcknowledgement.id).where(AlertAcknowledgement.user_id == user.id, AlertAcknowledgement.alert_key == alert_key)):
        db.add(AlertAcknowledgement(user_id=user.id, alert_key=alert_key[:160]))
        await db.commit()
    return {"acknowledged": True}


@router.get("/saved-filters")
async def saved_filters(db: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> dict:
    rows = (await db.scalars(select(SavedFilter).where(or_(SavedFilter.user_id == user.id, SavedFilter.is_shared.is_(True))).order_by(SavedFilter.name))).all()
    return {"items": [{"id": str(row.id), "name": row.name, "filters": row.filters, "is_shared": row.is_shared, "owner": str(row.user_id)} for row in rows]}


@router.post("/saved-filters", dependencies=[Depends(csrf_protect)])
async def save_filter(data: SavedFilterCreate, db: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> dict:
    if data.is_shared and user.role != "ADMIN":
        raise HTTPException(status_code=403, detail="Faqat admin umumiy filter yarata oladi")
    item = await db.scalar(select(SavedFilter).where(SavedFilter.user_id == user.id, SavedFilter.name == data.name))
    if item:
        item.filters = data.filters
        item.is_shared = data.is_shared
    else:
        item = SavedFilter(user_id=user.id, **data.model_dump())
        db.add(item)
    await db.commit()
    await db.refresh(item)
    return {"id": str(item.id), "name": item.name, "filters": item.filters, "is_shared": item.is_shared}


@router.delete("/saved-filters/{filter_id}", dependencies=[Depends(csrf_protect)])
async def delete_saved_filter(filter_id: uuid.UUID, db: AsyncSession = Depends(get_db), user: User = Depends(current_user)) -> Response:
    item = await db.get(SavedFilter, filter_id)
    if not item or (item.user_id != user.id and user.role != "ADMIN"):
        raise HTTPException(status_code=404, detail="Filter topilmadi")
    await db.delete(item)
    await db.commit()
    return Response(status_code=204)


@router.get("/search")
async def global_search(q: str = Query(min_length=2, max_length=80), db: AsyncSession = Depends(get_db)) -> dict:
    pattern = contains_pattern(q)
    posts = (await db.scalars(select(CustomsPost).where(CustomsPost.is_active.is_(True), or_(CustomsPost.post_code.ilike(pattern, escape="\\"), CustomsPost.post_name.ilike(pattern, escape="\\"))).limit(10))).all()
    corridors = (await db.scalars(select(Corridor).where(Corridor.is_active.is_(True), or_(Corridor.code.ilike(pattern, escape="\\"), Corridor.name.ilike(pattern, escape="\\"))).limit(10))).all()
    return {"posts": [{"id": str(item.id), "code": item.post_code, "name": item.post_name, "latitude": item.latitude, "longitude": item.longitude} for item in posts], "corridors": [{"id": str(item.id), "code": item.code, "name": item.name} for item in corridors]}


@router.get("/operations/backup")
async def backup(db: AsyncSession = Depends(get_db), _: User = Depends(admin_user)) -> Response:
    posts = (await db.scalars(select(CustomsPost).order_by(CustomsPost.post_code))).all()
    gateways = (await db.scalars(select(CountryGateway).order_by(CountryGateway.country_code, CountryGateway.name))).all()
    corridors = (await db.scalars(select(Corridor).options(selectinload(Corridor.waypoints)).order_by(Corridor.priority, Corridor.code))).unique().all()
    settings_rows = (await db.scalars(select(AppSetting))).all()
    payload = {
        "version": 1,
        "created_at": datetime.now(UTC),
        "posts": [{column.name: getattr(row, column.name) for column in CustomsPost.__table__.columns if column.name != "location"} for row in posts],
        "gateways": [{column.name: getattr(row, column.name) for column in CountryGateway.__table__.columns if column.name != "location"} for row in gateways],
        "corridors": [{
            "code": row.code, "name": row.name, "origin_country_code": row.origin_country_code, "destination_country_code": row.destination_country_code,
            "entry_post_code": row.entry_post_code, "exit_post_code": row.exit_post_code, "color": row.color,
            "routing_profile": row.routing_profile, "priority": row.priority, "is_active": row.is_active,
            "waypoints": [{"sequence_no": point.sequence_no, "waypoint_type": point.waypoint_type, "latitude": point.latitude, "longitude": point.longitude, "post_code": point.post_code, "gateway_id": str(point.gateway_id) if point.gateway_id else None, "label": point.label} for point in row.waypoints],
        } for row in corridors],
        "settings": [{"key": row.key, "value": row.value} for row in settings_rows],
        "note": "Katta deklaratsiya va metrikalar CSV import/eksport orqali ko'chiriladi.",
    }
    content = json.dumps(jsonable_encoder(payload), ensure_ascii=False, indent=2)
    return Response(content, media_type="application/json", headers={"Content-Disposition": f"attachment; filename=transport-backup-{date.today()}.json"})


@router.post("/operations/restore", dependencies=[Depends(csrf_protect)])
async def restore_backup(
    request: Request,
    background_tasks: BackgroundTasks,
    confirmation: str = Form(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(admin_user),
) -> dict:
    if confirmation != "RESTORE":
        raise HTTPException(status_code=422, detail="Tasdiqlash maydoniga RESTORE yozing")
    raw = await file.read(MAX_IMPORT_BYTES + 1)
    if len(raw) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Backup 5 MB dan katta bo'lmasin")
    try:
        payload = json.loads(raw.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=422, detail="Backup JSON fayli yaroqsiz") from exc
    if payload.get("version") != 1 or not isinstance(payload.get("posts"), list):
        raise HTTPException(status_code=422, detail="Backup versiyasi yoki tuzilmasi qo'llanmaydi")
    post_columns = {"post_code", "post_name", "post_type", "post_category", "region", "neighbor_country_code", "latitude", "longitude", "location_verified", "allow_passenger_vehicles", "allow_cargo_vehicles", "is_active"}
    post_rows = [{key: value for key, value in row.items() if key in post_columns} for row in payload["posts"] if row.get("post_code")]
    if post_rows:
        statement = insert(CustomsPost).values(post_rows)
        statement = statement.on_conflict_do_update(index_elements=[CustomsPost.post_code], set_={key: getattr(statement.excluded, key) for key in post_columns if key != "post_code"})
        await db.execute(statement)
        codes = [row["post_code"] for row in post_rows]
        await db.execute(update(CustomsPost).where(CustomsPost.post_code.in_(codes), CustomsPost.latitude.is_not(None), CustomsPost.longitude.is_not(None)).values(location=ST_SetSRID(ST_MakePoint(CustomsPost.longitude, CustomsPost.latitude), 4326)))
    gateway_columns = {"id", "country_code", "name", "gateway_type", "latitude", "longitude", "neighbor_country_code", "verified", "is_active", "notes"}
    gateway_rows = [{key: value for key, value in row.items() if key in gateway_columns} for row in payload.get("gateways", []) if row.get("id")]
    if gateway_rows:
        gateway_statement = insert(CountryGateway).values(gateway_rows)
        gateway_statement = gateway_statement.on_conflict_do_update(index_elements=[CountryGateway.id], set_={key: getattr(gateway_statement.excluded, key) for key in gateway_columns if key != "id"})
        await db.execute(gateway_statement)
        gateway_ids = [row["id"] for row in gateway_rows]
        await db.execute(update(CountryGateway).where(CountryGateway.id.in_(gateway_ids)).values(location=ST_SetSRID(ST_MakePoint(CountryGateway.longitude, CountryGateway.latitude), 4326)))
    restored_corridors: list[str] = []
    corridor_fields = {"name", "origin_country_code", "destination_country_code", "entry_post_code", "exit_post_code", "color", "routing_profile", "priority", "is_active"}
    for item in payload.get("corridors", []):
        if not item.get("code") or not item.get("waypoints"):
            continue
        corridor = await db.scalar(select(Corridor).where(Corridor.code == item["code"]))
        values = {key: item.get(key) for key in corridor_fields if key in item}
        if corridor is None:
            corridor = Corridor(code=item["code"], status="REVIEW", route_needs_review=True, geometry_source="backup-restore", **values)
            db.add(corridor)
            await db.flush()
        else:
            for key, value in values.items():
                setattr(corridor, key, value)
            corridor.geometry = None
            corridor.status = "REVIEW"
            corridor.route_needs_review = True
            corridor.geometry_source = "backup-restore"
            await db.execute(delete(CorridorWaypoint).where(CorridorWaypoint.corridor_id == corridor.id))
        for point in item["waypoints"]:
            db.add(waypoint_model(corridor.id, point))
        restored_corridors.append(str(corridor.id))
    for item in payload.get("settings", []):
        if not item.get("key") or not isinstance(item.get("value"), dict):
            continue
        setting = await db.get(AppSetting, item["key"])
        if setting:
            setting.value = item["value"]
            setting.updated_by = user.id
        else:
            db.add(AppSetting(key=item["key"], value=item["value"], updated_by=user.id))
    job = None
    if restored_corridors:
        job = BackgroundJob(kind="ROUTE_REBUILD", total=len(restored_corridors), payload={"corridor_ids": restored_corridors, "routing_profile": "driving", "reason": "backup_restore"}, created_by=user.id)
        db.add(job)
        await db.flush()
    await add_audit(db, request, user, "RESTORE_BACKUP", "backup", file.filename, after={"posts": len(post_rows), "corridors": len(restored_corridors), "job_id": str(job.id) if job else None})
    await db.commit()
    if job:
        background_tasks.add_task(run_route_rebuild, job.id)
    return {"posts": len(post_rows), "gateways": len(gateway_rows), "corridors": len(restored_corridors), "route_job_id": str(job.id) if job else None}
